"""Turn a validated upload into ExtractedContent.

* Digital PDFs: pdfplumber (pdfminer.six) for words and ruled tables.
* Scanned PDFs: pages rendered with pypdfium2, then OCR.
* Images: local Tesseract OCR via pytesseract.
* XLSX: openpyxl, values only (formulas are not evaluated or executed).
* CSV / TXT: standard library parsing.

Nothing here interprets meaning; see interpret.py.
"""

from __future__ import annotations

import csv
import io
import re
import shutil
import statistics
from collections.abc import Iterable, Sequence
from dataclasses import dataclass
from decimal import Decimal

from scholariq.ingestion.types import Block, Cell, ExtractedContent, Row, SourceKind
from scholariq.ingestion.validation import FileKind, ValidatedUpload

MAX_ROWS = 2000
# A PDF with fewer extractable characters than this is treated as scanned.
MIN_DIGITAL_CHARS = 40
OCR_DPI = 300


class ExtractionError(RuntimeError):
    """Extraction failed; the message is safe to show to the student."""


class OcrUnavailable(ExtractionError):
    pass


@dataclass(frozen=True)
class Word:
    text: str
    x0: float
    x1: float
    top: float
    bottom: float
    confidence: float = 1.0


def words_to_rows(words: Sequence[Word], line_key: Iterable[object] | None = None) -> list[Row]:
    """Group positioned words into rows and split each row into cells at wide gaps.

    Words on the same line closer than about one character height belong to
    the same cell; larger gaps start a new cell (column boundary).
    """
    if not words:
        return []
    heights = [w.bottom - w.top for w in words if w.bottom > w.top]
    unit = statistics.median(heights) if heights else 10.0

    if line_key is not None:
        groups: dict[object, list[Word]] = {}
        for key, w in zip(line_key, words, strict=True):
            groups.setdefault(key, []).append(w)
        lines = sorted(groups.values(), key=lambda ws: min(w.top for w in ws))
    else:
        lines = []
        for w in sorted(words, key=lambda w: (w.top, w.x0)):
            mid = (w.top + w.bottom) / 2
            if lines and abs(mid - _mid(lines[-1])) <= unit * 0.5:
                lines[-1].append(w)
            else:
                lines.append([w])

    rows: list[Row] = []
    for line in lines:
        line.sort(key=lambda w: w.x0)
        cells: list[list[Word]] = [[line[0]]]
        for prev, cur in zip(line, line[1:], strict=False):
            if cur.x0 - prev.x1 > unit * 1.0:
                cells.append([cur])
            else:
                cells[-1].append(cur)
        rows.append(
            Row(
                top=min(w.top for w in line),
                cells=[
                    Cell(
                        " ".join(w.text for w in c),
                        min(w.confidence for w in c),
                        x0=c[0].x0,
                        x1=c[-1].x1,
                    )
                    for c in cells
                ],
            )
        )
    return rows


def _mid(line: list[Word]) -> float:
    return statistics.fmean((w.top + w.bottom) / 2 for w in line)


# --------------------------------------------------------------------------- #
# PDF
# --------------------------------------------------------------------------- #


def _clean_cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return format(Decimal(repr(value)).normalize(), "f")
    return re.sub(r"\s+", " ", str(value)).strip()


def extract_pdf(data: bytes, tesseract_cmd: str | None = None) -> ExtractedContent:
    import pdfplumber

    blocks: list[Block] = []
    chars = 0
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        pages = len(pdf.pages)
        for number, page in enumerate(pdf.pages, 1):
            chars += len(page.chars)
            remaining = page
            # (top, table rows) for ruled tables; free text is placed around them.
            tables: list[tuple[float, list[Row]]] = []
            for table in page.find_tables():
                rows = [Row([Cell(_clean_cell(c)) for c in r]) for r in table.extract()]
                rows = [r for r in rows if not r.is_blank()]
                if rows:
                    tables.append((table.bbox[1], rows))
                remaining = remaining.outside_bbox(table.bbox)
            words = [
                Word(w["text"], w["x0"], w["x1"], w["top"], w["bottom"])
                for w in remaining.extract_words(x_tolerance=1.5, y_tolerance=2)
            ]
            blocks.extend(_order_page(words_to_rows(words), tables, number))

    if chars >= MIN_DIGITAL_CHARS:
        return ExtractedContent(SourceKind.PDF_TEXT, blocks, pages=pages)

    # Too little embedded text: treat as a scanned document.
    content = ExtractedContent(SourceKind.OCR, [], pages=pages)
    content.notes.append("No text layer found; pages were read with OCR.")
    for number, image in enumerate(_render_pdf(data), 1):
        for block in _ocr_image(image, tesseract_cmd).blocks:
            block.page = number
            content.blocks.append(block)
    return content


def _order_page(
    text_rows: list[Row], tables: list[tuple[float, list[Row]]], page: int
) -> list[Block]:
    """Interleave free-text rows and ruled tables in reading order.

    Text before a table often carries its context (e.g. "Semester III"), so the
    order matters to interpretation.
    """
    items: list[tuple[float, int, list[Row]]] = [(t, 1, rows) for t, rows in tables]
    items += [(r.top or 0.0, 0, [r]) for r in text_rows]
    items.sort(key=lambda it: (it[0], it[1]))
    blocks: list[Block] = []
    pending: list[Row] = []
    for _, is_table, rows in items:
        if is_table:
            if pending:
                blocks.append(Block(pending, page=page))
                pending = []
            blocks.append(Block(rows, page=page))
        else:
            pending.extend(rows)
    if pending:
        blocks.append(Block(pending, page=page))
    return blocks


def _render_pdf(data: bytes) -> Iterable[object]:
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(data)
    try:
        for page in pdf:
            yield page.render(scale=OCR_DPI / 72).to_pil()
    finally:
        pdf.close()


# --------------------------------------------------------------------------- #
# OCR
# --------------------------------------------------------------------------- #


def tesseract_available(tesseract_cmd: str | None = None) -> bool:
    return bool(tesseract_cmd or shutil.which("tesseract"))


def _preprocess(image: object) -> object:
    from PIL import Image, ImageOps

    assert isinstance(image, Image.Image)
    img = ImageOps.exif_transpose(image) or image
    img = img.convert("L")
    img = ImageOps.autocontrast(img, cutoff=1)
    if img.width < 1600:
        factor = 1600 / img.width
        img = img.resize(
            (int(img.width * factor), int(img.height * factor)), Image.Resampling.LANCZOS
        )
    return img


def _ocr_image(image: object, tesseract_cmd: str | None) -> ExtractedContent:
    import pytesseract

    if not tesseract_available(tesseract_cmd):
        raise OcrUnavailable(
            "This file needs OCR, but Tesseract is not installed on the server. "
            "Upload a digital PDF, spreadsheet or CSV instead, or ask the administrator "
            "to install Tesseract."
        )
    if tesseract_cmd:
        pytesseract.pytesseract.tesseract_cmd = tesseract_cmd

    img = _preprocess(image)
    # psm 6: assume a uniform block of text, which suits tabular marksheets.
    data = pytesseract.image_to_data(
        img, lang="eng", config="--psm 6", output_type=pytesseract.Output.DICT
    )
    words: list[Word] = []
    keys: list[tuple[int, int, int]] = []
    for i, text in enumerate(data["text"]):
        text = text.strip()
        conf = float(data["conf"][i])
        if not text or conf < 0:
            continue
        top, height = data["top"][i], data["height"][i]
        left, width = data["left"][i], data["width"][i]
        words.append(Word(text, left, left + width, top, top + height, conf / 100))
        keys.append((data["block_num"][i], data["par_num"][i], data["line_num"][i]))
    rows = words_to_rows(words, line_key=keys)
    return ExtractedContent(SourceKind.OCR, [Block(rows)] if rows else [])


def extract_image(data: bytes, tesseract_cmd: str | None = None) -> ExtractedContent:
    from PIL import Image

    with Image.open(io.BytesIO(data)) as img:
        img.load()
        content = _ocr_image(img, tesseract_cmd)
    content.notes.append("Read with OCR; please check every value.")
    return content


# --------------------------------------------------------------------------- #
# Spreadsheets and text
# --------------------------------------------------------------------------- #


def extract_xlsx(data: bytes) -> ExtractedContent:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    blocks: list[Block] = []
    try:
        for ws in wb.worksheets:
            rows: list[Row] = []
            for values in ws.iter_rows(values_only=True):
                row = Row([Cell(_clean_cell(v)) for v in values])
                if not row.is_blank():
                    rows.append(row)
                if len(rows) >= MAX_ROWS:
                    break
            if rows:
                blocks.append(Block(rows))
    finally:
        wb.close()
    return ExtractedContent(SourceKind.SPREADSHEET, blocks, pages=len(blocks))


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    raise ExtractionError("The text encoding could not be read.")


def extract_csv(data: bytes) -> ExtractedContent:
    text = _decode(data)
    try:
        dialect: type[csv.Dialect] | csv.Dialect = csv.Sniffer().sniff(
            text[:4096], delimiters=",;\t|"
        )
    except csv.Error:
        dialect = csv.excel
    rows = []
    for values in csv.reader(io.StringIO(text), dialect):
        row = Row([Cell(_clean_cell(v)) for v in values])
        if not row.is_blank():
            rows.append(row)
        if len(rows) >= MAX_ROWS:
            break
    return ExtractedContent(SourceKind.CSV, [Block(rows)] if rows else [])


_TEXT_SPLIT = re.compile(r"\t+|\s{2,}|\s*\|\s*")


def extract_text(data: bytes) -> ExtractedContent:
    rows = []
    for line in _decode(data).splitlines()[:MAX_ROWS]:
        parts = [p.strip() for p in _TEXT_SPLIT.split(line.strip()) if p.strip()]
        if parts:
            rows.append(Row([Cell(p) for p in parts]))
    return ExtractedContent(SourceKind.TEXT, [Block(rows)] if rows else [])


def extract(upload: ValidatedUpload, tesseract_cmd: str | None = None) -> ExtractedContent:
    try:
        match upload.kind:
            case FileKind.PDF:
                return extract_pdf(upload.data, tesseract_cmd)
            case FileKind.PNG | FileKind.JPEG:
                return extract_image(upload.data, tesseract_cmd)
            case FileKind.XLSX:
                return extract_xlsx(upload.data)
            case FileKind.CSV:
                return extract_csv(upload.data)
            case FileKind.TEXT:
                return extract_text(upload.data)
    except ExtractionError:
        raise
    except Exception as exc:  # noqa: BLE001 - parser libraries raise many types
        raise ExtractionError(
            f"The {upload.kind.value.upper()} file could not be read ({type(exc).__name__})."
        ) from exc
    raise ExtractionError("Unsupported file type.")
